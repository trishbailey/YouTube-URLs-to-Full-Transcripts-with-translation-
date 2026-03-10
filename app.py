import io
import re
from typing import Optional

import pandas as pd
import streamlit as st
from deep_translator import GoogleTranslator
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from youtube_transcript_api import YouTubeTranscriptApi
from youtube_transcript_api._errors import (
    NoTranscriptFound,
    TranscriptsDisabled,
    VideoUnavailable,
    YouTubeRequestFailed,
)


st.set_page_config(
    page_title="YouTube Transcript Extractor",
    page_icon="🎬",
    layout="wide",
)


LANG_CODE_MAP = {
    "None (keep original)": None,
    "Burmese (Myanmar) → English": "my",
    "Filipino / Tagalog → English": "tl",
    "Indonesian → English": "id",
    "Khmer (Cambodian) → English": "km",
    "Lao → English": "lo",
    "Malay → English": "ms",
    "Thai → English": "th",
    "Vietnamese → English": "vi",
    "Bengali → English": "bn",
    "Cebuano → English": "ceb",
    "Javanese → English": "jw",
    "Sundanese → English": "su",
    "Tetum (Timor-Leste) → English": "tet",
}

FALLBACK_LANGS = ["en", "id", "ms", "th", "vi", "tl"]
MAX_TRANSLATION_CHARS = 4500


def extract_video_id(value: str) -> Optional[str]:
    """Extract an 11-character YouTube video ID from a URL or raw ID."""
    if not value:
        return None

    value = value.strip()
    if re.fullmatch(r"[0-9A-Za-z_-]{11}", value):
        return value

    patterns = [
        r"(?:v=|\/)([0-9A-Za-z_-]{11})(?:[&?\/]|$)",
        r"youtu\.be\/([0-9A-Za-z_-]{11})(?:[?&\/]|$)",
        r"embed\/([0-9A-Za-z_-]{11})(?:[?&\/]|$)",
        r"shorts\/([0-9A-Za-z_-]{11})(?:[?&\/]|$)",
        r"live\/([0-9A-Za-z_-]{11})(?:[?&\/]|$)",
    ]
    for pattern in patterns:
        match = re.search(pattern, value)
        if match:
            return match.group(1)
    return None


@st.cache_resource(show_spinner=False)
def get_transcript_client() -> YouTubeTranscriptApi:
    return YouTubeTranscriptApi()


def chunk_text(text: str, max_chars: int = MAX_TRANSLATION_CHARS) -> list[str]:
    """Split text into chunks without breaking words where possible."""
    if len(text) <= max_chars:
        return [text]

    words = text.split()
    chunks: list[str] = []
    current: list[str] = []
    current_len = 0

    for word in words:
        proposed_len = current_len + len(word) + (1 if current else 0)
        if proposed_len > max_chars and current:
            chunks.append(" ".join(current))
            current = [word]
            current_len = len(word)
        else:
            current.append(word)
            current_len = proposed_len

    if current:
        chunks.append(" ".join(current))
    return chunks


def translate_to_english(text: str, source_lang_code: str) -> str:
    translator = GoogleTranslator(source=source_lang_code, target="en")
    translated_chunks = []
    for chunk in chunk_text(text):
        translated_chunks.append(translator.translate(chunk) or "")
    return " ".join(part.strip() for part in translated_chunks if part).strip()


def transcript_to_text(fetched_transcript) -> str:
    return " ".join(snippet.text.strip() for snippet in fetched_transcript if getattr(snippet, "text", "")).strip()


def choose_transcript(transcript_list, requested_lang: Optional[str]):
    """Select the best transcript available.

    Order:
    1. Exact requested language, manual or generated.
    2. Preferred fallback languages.
    3. First manual transcript.
    4. First generated transcript.
    """
    all_tracks = list(transcript_list)
    if not all_tracks:
        raise NoTranscriptFound("No transcript tracks available")

    if requested_lang:
        try:
            return transcript_list.find_transcript([requested_lang])
        except Exception:
            pass

    preferred = [lang for lang in FALLBACK_LANGS if lang != requested_lang]
    if preferred:
        try:
            return transcript_list.find_transcript(preferred)
        except Exception:
            pass

    manual_tracks = [t for t in all_tracks if not getattr(t, "is_generated", False)]
    if manual_tracks:
        return manual_tracks[0]

    return all_tracks[0]


def fetch_transcript(video_id: str, source_lang_code: Optional[str]) -> dict:
    result = {
        "video_id": video_id,
        "url": f"https://www.youtube.com/watch?v={video_id}",
        "language_detected": "",
        "language_label": "",
        "translated": False,
        "track_type": "",
        "transcript": "",
        "error": "",
    }

    try:
        client = get_transcript_client()
        transcript_list = client.list(video_id)
        transcript_obj = choose_transcript(transcript_list, source_lang_code)

        result["language_detected"] = transcript_obj.language_code
        result["language_label"] = getattr(transcript_obj, "language", "")
        result["track_type"] = "Auto-generated" if getattr(transcript_obj, "is_generated", False) else "Manual"

        fetched = transcript_obj.fetch()
        original_text = transcript_to_text(fetched)

        if not original_text:
            result["error"] = "Transcript track was found, but no text was returned."
            return result

        if source_lang_code and transcript_obj.language_code == source_lang_code:
            translated_text = translate_to_english(original_text, source_lang_code)
            if translated_text:
                result["transcript"] = translated_text
                result["translated"] = True
            else:
                result["transcript"] = original_text
        else:
            result["transcript"] = original_text

    except TranscriptsDisabled:
        result["error"] = "Transcripts are disabled for this video."
    except NoTranscriptFound:
        result["error"] = "No transcript could be found for this video."
    except VideoUnavailable:
        result["error"] = "Video is unavailable, private, removed, or region-restricted."
    except YouTubeRequestFailed as exc:
        result["error"] = f"YouTube request failed: {exc}"
    except Exception as exc:
        result["error"] = str(exc)

    return result


def build_excel(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Transcripts"

    header_fill = PatternFill("solid", start_color="1F4E79", end_color="1F4E79")
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    alt_fill = PatternFill("solid", start_color="DCE6F1", end_color="DCE6F1")
    white_fill = PatternFill("solid", start_color="FFFFFF", end_color="FFFFFF")

    cell_font = Font(name="Arial", size=10)
    cell_align = Alignment(vertical="top", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )

    headers = [
        "#",
        "YouTube URL",
        "Video ID",
        "Language Code",
        "Language",
        "Track Type",
        "Translated to EN?",
        "Status",
        "Transcript / Error",
    ]
    col_widths = [5, 46, 16, 16, 22, 16, 16, 14, 110]

    for idx, (header, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=1, column=idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    for row_num, row in enumerate(rows, start=2):
        fill = alt_fill if row_num % 2 == 0 else white_fill
        status = "Error" if row["error"] else "Success"
        transcript_or_error = row["error"] or row["transcript"]
        values = [
            row_num - 1,
            row["url"],
            row["video_id"],
            row["language_detected"],
            row["language_label"],
            row["track_type"],
            "Yes" if row["translated"] else "No",
            status,
            transcript_or_error,
        ]

        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_num, column=col_idx, value=value)
            cell.font = cell_font
            cell.fill = fill
            cell.alignment = cell_align
            cell.border = thin_border

        approx_height = max(20, min(400, len(str(transcript_or_error)) // 12))
        ws.row_dimensions[row_num].height = approx_height

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def normalize_uploaded_urls(df: pd.DataFrame) -> list[str]:
    cols_lower = [str(c).strip().lower() for c in df.columns]

    if "id" in cols_lower and "url" not in cols_lower:
        id_col = df.columns[cols_lower.index("id")]
        return [
            f"https://www.youtube.com/watch?v={video_id}"
            for video_id in df[id_col].dropna().astype(str).str.strip()
            if extract_video_id(video_id)
        ]

    for candidate in ["url", "youtube_url", "video_url", "link"]:
        if candidate in cols_lower:
            source_col = df.columns[cols_lower.index(candidate)]
            return df[source_col].dropna().astype(str).str.strip().tolist()

    raise ValueError("No supported column found. Use one of: id, url, youtube_url, video_url, link")


def dedupe_preserve_order(urls: list[str]) -> list[str]:
    seen = set()
    clean_urls = []
    for url in urls:
        url = url.strip()
        if not url or url in seen:
            continue
        seen.add(url)
        clean_urls.append(url)
    return clean_urls


st.title("🎬 YouTube Transcript Extractor")
st.caption("Paste YouTube URLs or upload a file, then export a formatted Excel workbook.")

with st.sidebar:
    st.header("Options")
    translation_option = st.selectbox(
        "Translate to English from…",
        options=list(LANG_CODE_MAP.keys()),
        help="Choose a source language only if you want English output for transcripts in that language.",
    )
    st.markdown("---")
    st.markdown(
        "**Accepted input**\n"
        "- Full YouTube URLs\n"
        "- youtu.be URLs\n"
        "- shorts URLs\n"
        "- bare 11-character video IDs\n"
        "- CSV/XLSX with `url` or `id` column"
    )
    st.markdown("---")
    st.markdown(
        "**Operational notes**\n"
        "- Videos must have captions available.\n"
        "- Some videos fail if YouTube blocks the request or the video is private.\n"
        "- Translation uses Google Translate via `deep-translator`."
    )

st.subheader("Input")
tab_paste, tab_file = st.tabs(["Paste URLs", "Upload CSV / Excel"])

raw_urls: list[str] = []

with tab_paste:
    url_input = st.text_area(
        "One URL or video ID per line",
        placeholder="https://www.youtube.com/watch?v=dQw4w9WgXcQ\ndQw4w9WgXcQ",
        height=220,
    )
    if url_input.strip():
        raw_urls = [line.strip() for line in url_input.splitlines() if line.strip()]

with tab_file:
    uploaded_file = st.file_uploader("Upload CSV or Excel", type=["csv", "xlsx", "xls"])
    if uploaded_file is not None:
        try:
            file_name = uploaded_file.name.lower()
            if file_name.endswith((".xlsx", ".xls")):
                df_upload = pd.read_excel(uploaded_file)
            else:
                df_upload = pd.read_csv(uploaded_file)

            raw_urls = normalize_uploaded_urls(df_upload)
            st.success(f"Loaded {len(raw_urls)} rows from {uploaded_file.name}.")
            with st.expander("Preview input rows"):
                st.dataframe(pd.DataFrame({"input": raw_urls[:25]}), use_container_width=True)
        except Exception as exc:
            st.error(f"Could not read file: {exc}")

raw_urls = dedupe_preserve_order(raw_urls)

if raw_urls:
    st.info(f"Ready to process {len(raw_urls)} unique input rows.")

run = st.button(
    f"Extract transcripts{f' ({len(raw_urls)})' if raw_urls else ''}",
    type="primary",
    use_container_width=True,
    disabled=not raw_urls,
)

if run:
    requested_lang = LANG_CODE_MAP[translation_option]
    results: list[dict] = []

    progress = st.progress(0, text="Starting extraction…")
    status_box = st.empty()

    for idx, raw_value in enumerate(raw_urls, start=1):
        status_box.info(f"Processing {idx}/{len(raw_urls)}: {raw_value}")
        video_id = extract_video_id(raw_value)

        if not video_id:
            results.append(
                {
                    "video_id": "INVALID",
                    "url": raw_value,
                    "language_detected": "",
                    "language_label": "",
                    "track_type": "",
                    "translated": False,
                    "transcript": "",
                    "error": "Could not parse a valid YouTube video ID from this input.",
                }
            )
        else:
            results.append(fetch_transcript(video_id, requested_lang))

        progress.progress(idx / len(raw_urls), text=f"Processed {idx} of {len(raw_urls)}")

    progress.empty()
    status_box.empty()

    st.session_state["results"] = results
    st.session_state["excel_bytes"] = build_excel(results)

if "results" in st.session_state:
    results = st.session_state["results"]
    successes = sum(1 for row in results if not row["error"])
    failures = len(results) - successes

    a, b, c = st.columns(3)
    a.metric("Total inputs", len(results))
    b.metric("Succeeded", successes)
    c.metric("Failed", failures)

    if failures:
        with st.expander("Failures"):
            for row in results:
                if row["error"]:
                    st.error(f"{row['url']} — {row['error']}")

    preview_rows = []
    for row in results:
        preview_rows.append(
            {
                "url": row["url"],
                "video_id": row["video_id"],
                "language": row["language_detected"],
                "track_type": row["track_type"],
                "translated": row["translated"],
                "status": "Error" if row["error"] else "Success",
            }
        )

    st.subheader("Run summary")
    st.dataframe(pd.DataFrame(preview_rows), use_container_width=True)

    with st.expander("Transcript preview"):
        for row in results:
            if row["error"]:
                continue
            label = f"{row['url']} | {row['language_detected']} | {row['track_type']}"
            with st.expander(label):
                text = row["transcript"]
                st.write(text[:4000] + ("…" if len(text) > 4000 else ""))

    st.download_button(
        label="Download Excel",
        data=st.session_state["excel_bytes"],
        file_name="youtube_transcripts.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
        use_container_width=True,
    )
